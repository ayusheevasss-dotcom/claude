import openpyxl,glob,json,pandas as pd
U='/root/.claude/uploads/178766d5-eb4a-5539-b620-ef59ec9f56fb/'
mz={}
for f in glob.glob(U+'*-2026-*.xlsx'):
    m=int(f.split('-2026-')[1][:2])
    wb=openpyxl.load_workbook(f,data_only=True)
    ws=wb.worksheets[-1] if m!=1 else wb['31 января']
    rows=list(ws.iter_rows(values_only=True))
    tt = rows[1][2]=='т.т'
    def get(name):
        for r in rows:
            if r[0] and str(r[0]).strip().startswith(name): return r
    d={}
    for k,n in [('gross','Гросс'),('tab','Таборный'),('vrez','Врезанный')]:
        r=get(n); d[k+'_m3']=r[1] or 0; d[k+'_t']=(r[2] if tt else None); d[k+'_ouf_m3']=(r[3] if tt else r[2]) or 0
    i=[j for j,r in enumerate(rows) if r[0] and str(r[0]).startswith('Укладка КВ Гросс')][0]
    tc=7 if tt else 4
    v=0;
    for r in rows[i+3:]:
        if r[0] and str(r[0]).strip()=='Всего:': d['kv_t']=r[tc]; break
        if isinstance(r[1],(int,float)) and r[0] not in ('Автоукладка','Стакерная укладка'): v+=r[1]
    d['kv_m3']=v
    mz[m]=d
# Tokko
wb=openpyxl.load_workbook(glob.glob(U+'a13bba40*')[0],data_only=True); ws=wb['Сентябрь 2026']
rows=list(ws.iter_rows(values_only=True))
st={0:(5,6),1:(19,20),2:(32,33),3:(45,46)}
for m in range(1,10):
    c=6+2*(m-1)
    for s,(rt,rv) in st.items():
        mz[m][f'tokko{s}_t']=rows[rt][c]; mz[m][f'tokko{s}_m3']=rows[rv][c]
# Лист3 stages
for f,m in [('4e17a8a3-2026-07-31.xlsx',7),('7ae66ba5-2026-08-31.xlsx',8)]:
    ws=openpyxl.load_workbook(U+f,data_only=True)['Лист3']; r=list(ws.iter_rows(values_only=True))
    mz[m]['tabst']={'WST4':r[5][2],'EST3':r[5][3],'EST4':r[5][4]}
    mz[m]['grst']={'ST3':r[12][1],'ST4':r[12][2],'ST5':r[12][3],'ST6':r[12][4]}
d=pd.read_excel(glob.glob(U+'60308d46*')[0]); d.columns=['y','m','pit','pl','mat','wet','dry','vol']
for c in ['wet','dry','vol']: d[c]/=1000
ou={}
for m in range(1,10):
    x=d[d.m==m]; o={}
    def s(q): return float(q.wet.sum()),float(q.vol.sum())
    o['gross']=s(x[x.pit=='Gross']); o['kv']=s(x[x.pit=='MK Gross GL']); o['tab']=s(x[x.pit=='Tabornoe']); o['tokko']=s(x[x.pit=='Tokko'])
    o['vrez']=s(x[(x.pit=='Tabornoe')&(x.pl=='VREZ')])
    for p in ['ST2','ST3','ST4','ST5','ST6','ST7','Осыпи Гросс']: o['g_'+p]=s(x[(x.pit=='Gross')&(x.pl==p)])
    for p in ['EST1','EST3','EST4','WST4','VREZ','Sklad podsypka Tab']: o['t_'+p]=s(x[(x.pit=='Tabornoe')&(x.pl==p)])
    for p in ['ST-1','ST-2','ST-3']: o['k_'+p]=s(x[(x.pit=='Tokko')&(x.pl==p)])
    o['k_other']=s(x[(x.pit=='Tokko')&~x.pl.isin(['ST-1','ST-2','ST-3'])])
    ou[m]=o
json.dump({'mz':mz,'ou':ou},open('data.json','w'),ensure_ascii=False,indent=1,default=str)
for m in range(1,10): print(m,{k:(round(v,1) if isinstance(v,(int,float)) else v) for k,v in mz[m].items()})
for m in range(1,10): print(m,{k:(round(v[0],1),round(v[1],1)) for k,v in ou[m].items() if v[0]})
